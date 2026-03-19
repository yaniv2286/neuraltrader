#!/usr/bin/env python3
"""
Training Data Validation Script
Audits all training data sources and validates data quality
"""

import os
import sys
import json
import logging
from pathlib import Path
from datetime import datetime
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s [%(levelname)s] %(message)s'
)
logger = logging.getLogger('TrainingDataValidator')

DATA_DIR = PROJECT_ROOT / 'data' / 'raw'
MODELS_DIR = PROJECT_ROOT / 'models'
OUTPUT_DIR = PROJECT_ROOT / 'reports' / 'validation'


def audit_parquet_files():
    """Audit all parquet files in data/raw"""
    logger.info(f"[AUDIT] Scanning {DATA_DIR} for parquet files...")
    
    parquet_files = sorted(DATA_DIR.glob('*.parquet'))
    logger.info(f"[AUDIT] Found {len(parquet_files)} parquet files")
    
    inventory = []
    total_rows = 0
    
    for i, pf in enumerate(parquet_files, 1):
        ticker = pf.stem.upper()
        try:
            df = pd.read_parquet(pf)
            n_rows = len(df)
            total_rows += n_rows
            
            if 'date' in df.columns:
                df['date'] = pd.to_datetime(df['date'])
                start_date = df['date'].min()
                end_date = df['date'].max()
            elif df.index.name == 'date' or isinstance(df.index, pd.DatetimeIndex):
                start_date = df.index.min()
                end_date = df.index.max()
            else:
                start_date = end_date = None
            
            inventory.append({
                'ticker': ticker,
                'rows': n_rows,
                'start_date': start_date,
                'end_date': end_date,
                'file_size_mb': pf.stat().st_size / (1024 * 1024),
                'status': 'OK'
            })
            
            if i % 200 == 0:
                logger.info(f"[AUDIT] Processed {i}/{len(parquet_files)} files...")
                
        except Exception as e:
            inventory.append({
                'ticker': ticker,
                'rows': 0,
                'start_date': None,
                'end_date': None,
                'file_size_mb': 0,
                'status': f'ERROR: {str(e)[:50]}'
            })
    
    logger.info(f"[AUDIT] Total rows across all tickers: {total_rows:,}")
    return pd.DataFrame(inventory), total_rows


def document_feature_pipeline():
    """Document the 68-feature engineering pipeline"""
    features = [
        # Raw OHLCV (12 features)
        {'name': 'Price', 'category': 'Raw', 'formula': 'Close price'},
        {'name': 'high', 'category': 'Raw', 'formula': 'High price'},
        {'name': 'low', 'category': 'Raw', 'formula': 'Low price'},
        {'name': 'open', 'category': 'Raw', 'formula': 'Open price'},
        {'name': 'volume', 'category': 'Raw', 'formula': 'Trading volume'},
        {'name': 'Adj Close', 'category': 'Raw', 'formula': 'Adjusted close'},
        {'name': 'Adj High', 'category': 'Raw', 'formula': 'Adjusted high'},
        {'name': 'Adj Low', 'category': 'Raw', 'formula': 'Adjusted low'},
        {'name': 'Adj Open', 'category': 'Raw', 'formula': 'Adjusted open'},
        {'name': 'Adj Volume', 'category': 'Raw', 'formula': 'Adjusted volume'},
        {'name': 'Dividend', 'category': 'Raw', 'formula': 'Dividend amount'},
        {'name': 'Split Factor', 'category': 'Raw', 'formula': 'Stock split factor'},
        
        # Technical Indicators (52 features)
        {'name': 'sma_20', 'category': 'Technical', 'formula': '20-day simple moving average'},
        {'name': 'ema_20', 'category': 'Technical', 'formula': '20-day exponential moving average'},
        {'name': 'rsi', 'category': 'Technical', 'formula': 'Relative Strength Index (14-day)'},
        {'name': 'obv', 'category': 'Technical', 'formula': 'On-Balance Volume'},
        {'name': 'vwap', 'category': 'Technical', 'formula': 'Volume Weighted Average Price'},
        {'name': 'rolling_volatility', 'category': 'Technical', 'formula': '20-day rolling std of returns'},
        {'name': 'atr_14', 'category': 'Technical', 'formula': 'Average True Range (14-day)'},
        {'name': 'momentum', 'category': 'Technical', 'formula': 'Price momentum'},
        {'name': 'roc', 'category': 'Technical', 'formula': 'Rate of Change'},
        {'name': 'body_size', 'category': 'Technical', 'formula': 'Candle body size (close - open)'},
        {'name': 'upper_wick', 'category': 'Technical', 'formula': 'Upper wick (high - max(open, close))'},
        {'name': 'lower_wick', 'category': 'Technical', 'formula': 'Lower wick (min(open, close) - low)'},
        {'name': 'drawdown_pct', 'category': 'Technical', 'formula': 'Drawdown from peak'},
        {'name': 'close_lag_1', 'category': 'Technical', 'formula': 'Previous day close'},
        {'name': 'volume_lag_1', 'category': 'Technical', 'formula': 'Previous day volume'},
        {'name': 'rsi_lag_1', 'category': 'Technical', 'formula': 'Previous day RSI'},
        {'name': 'momentum_5', 'category': 'Technical', 'formula': '5-day momentum'},
        {'name': 'momentum_10', 'category': 'Technical', 'formula': '10-day momentum'},
        {'name': 'momentum_20', 'category': 'Technical', 'formula': '20-day momentum'},
        {'name': 'roc_5', 'category': 'Technical', 'formula': '5-day rate of change'},
        {'name': 'roc_10', 'category': 'Technical', 'formula': '10-day rate of change'},
        {'name': 'roc_20', 'category': 'Technical', 'formula': '20-day rate of change'},
        {'name': 'volume_ratio', 'category': 'Technical', 'formula': 'Volume / avg volume'},
        {'name': 'volume_log', 'category': 'Technical', 'formula': 'Log of volume'},
        {'name': 'macd', 'category': 'Technical', 'formula': 'MACD line'},
        {'name': 'macd_signal', 'category': 'Technical', 'formula': 'MACD signal line'},
        {'name': 'macd_histogram', 'category': 'Technical', 'formula': 'MACD histogram'},
        {'name': 'obv_ratio', 'category': 'Technical', 'formula': 'OBV ratio'},
        {'name': 'volatility_20', 'category': 'Technical', 'formula': '20-day volatility'},
        {'name': 'volatility_50', 'category': 'Technical', 'formula': '50-day volatility'},
        {'name': 'atr_ratio', 'category': 'Technical', 'formula': 'ATR ratio'},
        {'name': 'price_efficiency', 'category': 'Technical', 'formula': 'Price efficiency ratio'},
        {'name': 'bb_upper', 'category': 'Technical', 'formula': 'Bollinger Band upper'},
        {'name': 'bb_lower', 'category': 'Technical', 'formula': 'Bollinger Band lower'},
        {'name': 'bb_middle', 'category': 'Technical', 'formula': 'Bollinger Band middle'},
        {'name': 'bb_width', 'category': 'Technical', 'formula': 'Bollinger Band width'},
        {'name': 'bb_position', 'category': 'Technical', 'formula': 'Price position in BB'},
        {'name': 'bb_overbought', 'category': 'Technical', 'formula': 'BB overbought flag'},
        {'name': 'bb_oversold', 'category': 'Technical', 'formula': 'BB oversold flag'},
        {'name': 'vol_regime', 'category': 'Technical', 'formula': 'Volatility regime'},
        {'name': 'high_vol', 'category': 'Technical', 'formula': 'High volatility flag'},
        {'name': 'low_vol', 'category': 'Technical', 'formula': 'Low volatility flag'},
        {'name': 'trend_regime', 'category': 'Technical', 'formula': 'Trend regime'},
        {'name': 'strong_uptrend', 'category': 'Technical', 'formula': 'Strong uptrend flag'},
        {'name': 'strong_downtrend', 'category': 'Technical', 'formula': 'Strong downtrend flag'},
        {'name': 'sma10_sma50_cross', 'category': 'Technical', 'formula': 'SMA10/SMA50 crossover'},
        {'name': 'sma50_sma200_cross', 'category': 'Technical', 'formula': 'SMA50/SMA200 crossover'},
        {'name': 'macd_bullish', 'category': 'Technical', 'formula': 'MACD bullish signal'},
        {'name': 'macd_bearish', 'category': 'Technical', 'formula': 'MACD bearish signal'},
        {'name': 'price_sma10_ratio', 'category': 'Technical', 'formula': 'Price / SMA10'},
        {'name': 'price_sma50_ratio', 'category': 'Technical', 'formula': 'Price / SMA50'},
        {'name': 'price_sma200_ratio', 'category': 'Technical', 'formula': 'Price / SMA200'},
        
        # Phase 12 Momentum Features (4 features)
        {'name': 'rs_vs_spy', 'category': 'Momentum', 'formula': '20-day return vs SPY'},
        {'name': 'roc_63', 'category': 'Momentum', 'formula': '63-day (quarter) rate of change'},
        {'name': 'high_52w_prox', 'category': 'Momentum', 'formula': 'Close / 52-week high'},
        {'name': 'volume_breakout', 'category': 'Momentum', 'formula': 'Volume / 50-day avg volume'},
    ]
    
    return pd.DataFrame(features)


def analyze_data_quality(inventory_df):
    """Analyze data quality metrics"""
    quality_metrics = {
        'total_tickers': len(inventory_df),
        'valid_tickers': len(inventory_df[inventory_df['status'] == 'OK']),
        'error_tickers': len(inventory_df[inventory_df['status'] != 'OK']),
        'total_rows': inventory_df['rows'].sum(),
        'avg_rows_per_ticker': inventory_df['rows'].mean(),
        'min_rows': inventory_df['rows'].min(),
        'max_rows': inventory_df['rows'].max(),
        'total_size_gb': inventory_df['file_size_mb'].sum() / 1024,
    }
    
    date_range = inventory_df[inventory_df['start_date'].notna()]
    if len(date_range) > 0:
        quality_metrics['earliest_date'] = date_range['start_date'].min()
        quality_metrics['latest_date'] = date_range['end_date'].max()
    
    return quality_metrics


def document_sentiment_history():
    """Document the Phase 10 sentiment experiment"""
    history = [
        {
            'phase': 'Phase 10',
            'date': '2026-02-20',
            'description': 'Sentiment Integration',
            'features': 76,
            'breakdown': '64 technical + 12 sentiment',
            'status': 'COMPLETE'
        },
        {
            'phase': 'Phase 10.1',
            'date': '2026-02-21',
            'description': 'Model Retraining with Sentiment',
            'features': 116,
            'breakdown': '64 technical + 52 sentiment',
            'accuracy': '53.5-53.8%',
            'status': 'COMPLETE - Performance degraded'
        },
        {
            'phase': 'Phase 12',
            'date': '2026-03-14',
            'description': 'Pure Technical Models',
            'features': 68,
            'breakdown': '64 technical + 4 momentum',
            'accuracy': '68.5% @ 0.50, 91.4% precision @ 0.65',
            'status': 'CURRENT - Sentiment removed'
        }
    ]
    
    return pd.DataFrame(history)


def create_excel_report(inventory_df, features_df, quality_metrics, history_df, output_path):
    """Create comprehensive Excel report"""
    logger.info(f"[EXCEL] Creating report at {output_path}")
    
    wb = Workbook()
    wb.remove(wb.active)
    
    # Sheet 1: Ticker Inventory
    ws1 = wb.create_sheet("Ticker Inventory")
    ws1.append(['Ticker', 'Rows', 'Start Date', 'End Date', 'Size (MB)', 'Status'])
    
    for _, row in inventory_df.iterrows():
        ws1.append([
            row['ticker'],
            row['rows'],
            row['start_date'].strftime('%Y-%m-%d') if pd.notna(row['start_date']) else '',
            row['end_date'].strftime('%Y-%m-%d') if pd.notna(row['end_date']) else '',
            round(row['file_size_mb'], 2),
            row['status']
        ])
    
    # Header formatting
    for cell in ws1[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        cell.font = Font(bold=True, color='FFFFFF')
    
    # Sheet 2: Feature Pipeline
    ws2 = wb.create_sheet("Feature Pipeline")
    ws2.append(['Feature Name', 'Category', 'Formula/Description'])
    
    for _, row in features_df.iterrows():
        ws2.append([row['name'], row['category'], row['formula']])
    
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        cell.font = Font(bold=True, color='FFFFFF')
    
    # Sheet 3: Data Quality Metrics
    ws3 = wb.create_sheet("Data Quality")
    ws3.append(['Metric', 'Value'])
    
    for key, value in quality_metrics.items():
        if isinstance(value, (int, float)):
            if key.endswith('_gb'):
                ws3.append([key.replace('_', ' ').title(), f"{value:.2f} GB"])
            elif isinstance(value, float):
                ws3.append([key.replace('_', ' ').title(), f"{value:,.2f}"])
            else:
                ws3.append([key.replace('_', ' ').title(), f"{value:,}"])
        else:
            ws3.append([key.replace('_', ' ').title(), str(value)])
    
    for cell in ws3[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        cell.font = Font(bold=True, color='FFFFFF')
    
    # Sheet 4: Sentiment Experiment History
    ws4 = wb.create_sheet("Sentiment History")
    ws4.append(['Phase', 'Date', 'Description', 'Features', 'Breakdown', 'Accuracy', 'Status'])
    
    for _, row in history_df.iterrows():
        ws4.append([
            row['phase'],
            row['date'],
            row['description'],
            row['features'],
            row['breakdown'],
            row.get('accuracy', ''),
            row['status']
        ])
    
    for cell in ws4[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        cell.font = Font(bold=True, color='FFFFFF')
    
    # Adjust column widths
    for ws in [ws1, ws2, ws3, ws4]:
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    wb.save(output_path)
    logger.info(f"[EXCEL] Report saved successfully")


def main():
    logger.info("=" * 80)
    logger.info("TRAINING DATA VALIDATION - Part 5")
    logger.info("=" * 80)
    
    # Step 1: Audit parquet files
    inventory_df, total_rows = audit_parquet_files()
    
    # Step 2: Document feature pipeline
    features_df = document_feature_pipeline()
    logger.info(f"[FEATURES] Documented {len(features_df)} features")
    
    # Step 3: Analyze data quality
    quality_metrics = analyze_data_quality(inventory_df)
    logger.info(f"[QUALITY] Valid tickers: {quality_metrics['valid_tickers']}/{quality_metrics['total_tickers']}")
    logger.info(f"[QUALITY] Total rows: {quality_metrics['total_rows']:,}")
    
    # Step 4: Document sentiment history
    history_df = document_sentiment_history()
    
    # Step 5: Create Excel report
    output_path = OUTPUT_DIR / 'training_data_audit.xlsx'
    create_excel_report(inventory_df, features_df, quality_metrics, history_df, output_path)
    
    # Summary
    logger.info("=" * 80)
    logger.info("VALIDATION SUMMARY")
    logger.info("=" * 80)
    logger.info(f"Total Tickers: {quality_metrics['total_tickers']}")
    logger.info(f"Valid Tickers: {quality_metrics['valid_tickers']}")
    logger.info(f"Total Rows: {quality_metrics['total_rows']:,}")
    logger.info(f"Features: {len(features_df)} (12 Raw + 52 Technical + 4 Momentum)")
    logger.info(f"Report: {output_path}")
    logger.info("=" * 80)
    
    return 0


if __name__ == '__main__':
    sys.exit(main())
